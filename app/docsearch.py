"""Private document search (RAG) for FRIDAY: Thai + English, CPU-only, no server.

build_index(docs_dir, index_dir) -> stats      incremental by file mtime/size
search(query, k=3, index_dir=...) -> [{file, text, score}]
search_docs(query) -> str                      LLM tool wrapper (uses DOCS_DIR / DOCS_INDEX)
release()                                      drop the embedding model from RAM now

Backend: DOCS_EMBED env, "fastembed:<hf model>" (default, in-process ONNX on CPU, never touches Ollama/VRAM)
or "ollama:<model>" (NOTE: with OLLAMA_MAX_LOADED_MODELS=1 this evicts the chat model).
The embedder is loaded on first use and freed after IDLE_S seconds without use.
"""
import json, os, re, threading, time
import numpy as np

ROOT = os.environ.get("FRIDAY_ROOT") or os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # project root; code lives in app/
DOCS_DIR = os.environ.get("DOCS_DIR", os.environ.get("FRIDAY_VAULT", os.path.join(ROOT, "data", "vault")))  # the Obsidian vault
os.environ.setdefault("FASTEMBED_CACHE_PATH", os.path.join(ROOT, "models", "fastembed"))  # default %TEMP% gets wiped by cleanup
DOCS_INDEX = os.environ.get("DOCS_INDEX", os.path.join(DOCS_DIR, ".index"))
EMBED = os.environ.get("DOCS_EMBED", "fastembed:sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2")
OLLAMA = os.environ.get("FRIDAY_OLLAMA", "http://127.0.0.1:11434")
EXTS = (".txt", ".md", ".pdf", ".docx")
# characters; Thai has no word spaces, so chunk by chars snapped to line/space breaks.
# 300 Thai chars ~ 120 tokens: stays under MiniLM's 128-token limit (measured 0.25-0.4 tok/char)
CHUNK, OVERLAP = 300, 60
IDLE_S = 120


# ---------- parsing + chunking ----------
def read_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
        return "\n".join([p.text for p in d.paragraphs] + cells)
    with open(path, "rb") as f:
        b = f.read()
    try:
        return b.decode("utf-8-sig")
    except UnicodeDecodeError:  # older Thai Windows text files are cp874/TIS-620, not UTF-8
        return b.decode("cp874", errors="replace")


def chunk(text, size=CHUNK, overlap=OVERLAP):
    text = re.sub(r"[ \t]+", " ", re.sub(r"\n\s*\n+", "\n\n", text)).strip()
    out, i = [], 0
    while i < len(text):
        end = min(i + size, len(text))
        if end < len(text):  # prefer to cut at paragraph > line > space, but not in the first half
            for sep in ("\n\n", "\n", " "):
                j = text.rfind(sep, i + size // 2, end)
                if j > 0:
                    end = j
                    break
        piece = text[i:end].strip()
        if piece:
            out.append(piece)
        if end >= len(text):
            break
        i = max(end - overlap, i + 1)
    return out


# ---------- embedding backends ----------
_model, _model_name, _timer, _lock = None, None, None, threading.Lock()


def _prefix(kind, name):
    if "e5" in name:  # multilingual-e5-large (fastembed built-in) expects these
        return f"{kind}: "
    if "qwen3-embedding" in name and kind == "query":
        return "Instruct: Given a question, retrieve passages that answer it\nQuery: "
    return ""


def embed(texts, kind="passage", backend=None):
    """kind: 'query' or 'passage'. Returns L2-normalised float32 [n, dim]."""
    global _model, _model_name, _timer
    backend = backend or EMBED
    kind_, name = backend.split(":", 1)
    texts = [_prefix(kind, name) + t for t in texts]
    if kind_ == "ollama":
        import requests
        r = requests.post(f"{OLLAMA}/api/embed", json={"model": name, "input": texts, "keep_alive": 0}, timeout=600)
        r.raise_for_status()
        v = np.asarray(r.json()["embeddings"], np.float32)
    else:
        with _lock:
            if _model_name != name:
                from fastembed import TextEmbedding
                import onnxruntime as ort
                # ORT worker threads busy-spin ~0.5 s after every run (measured: 4 cores at 100% after a 4 ms query).
                # fastembed exposes no knob for it, so patch SessionOptions just while the session is built.
                so_cls = ort.SessionOptions
                def _quiet_so():
                    so = so_cls()
                    so.add_session_config_entry("session.intra_op.allow_spinning", "0")
                    so.add_session_config_entry("session.inter_op.allow_spinning", "0")
                    return so
                ort.SessionOptions = _quiet_so
                try:
                    _model = TextEmbedding(name, threads=max(1, (os.cpu_count() or 4) // 4))  # leave most cores to the user
                finally:
                    ort.SessionOptions = so_cls
                _model_name = name
            v = np.asarray(list(_model.embed(texts, batch_size=16)), np.float32)
            if _timer:
                _timer.cancel()
            _timer = threading.Timer(IDLE_S, release)
            _timer.daemon = True
            _timer.start()
    return v / np.maximum(np.linalg.norm(v, axis=1, keepdims=True), 1e-9)


def release():
    global _model, _model_name
    with _lock:
        _model, _model_name = None, None
    import gc; gc.collect()


# ---------- index: emb.npy + meta.json ----------
def _load(index_dir):
    try:
        with open(os.path.join(index_dir, "meta.json"), encoding="utf-8") as f:
            meta = json.load(f)
        emb = np.load(os.path.join(index_dir, "emb.npy"))
    except (OSError, ValueError):
        return None, None
    if len(emb) != len(meta.get("chunks", [])):  # half-written index (crash between the two saves) -> rebuild
        return None, None
    return meta, emb


def build_index(docs_dir=DOCS_DIR, index_dir=DOCS_INDEX, backend=None, max_depth=None, max_bytes=None, max_files=None,
                skip=None):
    """Caps (§65, folders FRIDAY doesn't own): folder depth, file size, number of files — the rest is skipped.
    skip(path) -> True: never read (a deny rule); dropped from the index if it was in."""
    backend = backend or EMBED
    t0 = time.time()
    meta, emb = _load(index_dir)
    if not meta or meta.get("backend") != backend:
        meta, emb = None, None
    old = {c["file"]: [] for c in (meta or {"chunks": []})["chunks"]}
    for i, c in enumerate((meta or {"chunks": []})["chunks"]):
        old[c["file"]].append(i)
    stamps = (meta or {}).get("files", {})
    chunks, vecs, files, todo, errors = [], [], {}, [], {}
    for root, dirs, names in os.walk(docs_dir):
        dirs[:] = [d for d in dirs if not d.startswith(".") and d != "node_modules"]  # skip the index dir itself
        if max_depth is not None and root[len(docs_dir):].count(os.sep) >= max_depth:
            dirs.clear()
        for n in sorted(names):
            if not n.lower().endswith(EXTS) or n.startswith("~$") or (max_files and len(files) >= max_files):
                continue
            p = os.path.join(root, n)
            rel = os.path.relpath(p, docs_dir)
            try:
                st = os.stat(p)
            except OSError:  # renamed/deleted between walk and stat (Obsidian save, memory._write)
                continue
            if (max_bytes and st.st_size > max_bytes) or (skip and skip(p)):
                continue
            stamp = [st.st_mtime, st.st_size]
            files[rel] = stamp
            if stamps.get(rel) == stamp and rel in old:  # unchanged: reuse vectors
                for i in old[rel]:
                    chunks.append(meta["chunks"][i])
                    vecs.append(emb[i])
                continue
            try:
                todo += [{"file": rel, "text": t} for t in chunk(read_text(p))]
            except Exception as e:  # one broken file must not stop the index
                errors[rel] = f"{type(e).__name__}: {e}"
    if todo:
        new = embed([c["text"] for c in todo], "passage", backend)
        chunks += todo
        vecs += list(new)
    if not todo and meta and files == stamps:  # runs every chat turn: don't rewrite an unchanged index
        return {"files": len(files), "chunks": len(chunks), "embedded": 0, "errors": errors, "seconds": round(time.time() - t0, 2)}
    dim = len(vecs[0]) if vecs else 0
    emb = np.asarray(vecs, np.float32).reshape(len(vecs), dim)
    os.makedirs(index_dir, exist_ok=True)
    np.save(os.path.join(index_dir, "emb.tmp.npy"), emb)
    with open(os.path.join(index_dir, "meta.tmp"), "w", encoding="utf-8") as f:
        json.dump({"backend": backend, "files": files, "chunks": chunks}, f, ensure_ascii=False)
    os.replace(os.path.join(index_dir, "emb.tmp.npy"), os.path.join(index_dir, "emb.npy"))
    os.replace(os.path.join(index_dir, "meta.tmp"), os.path.join(index_dir, "meta.json"))  # meta last; _load checks they match
    _cache.clear()
    return {"files": len(files), "chunks": len(chunks), "embedded": len(todo), "errors": errors, "seconds": round(time.time() - t0, 2)}


_cache = {}


def search(query, k=3, index_dir=DOCS_INDEX):
    key = os.path.join(index_dir, "meta.json")
    try:
        mt = os.path.getmtime(key)
    except OSError:
        return []
    if _cache.get(key, (None,))[0] != mt:  # one entry per index: §65 searches 3-4 folder indexes, then the vault's
        _cache[key] = (mt, *_load(index_dir))
    _, meta, emb = _cache[key]
    if emb is None or not len(emb):
        return []
    s = emb @ embed([query], "query", meta["backend"])[0]
    out, seen = [], {}
    for i in np.argsort(-s):  # max 2 passages per file so k results cover more documents
        f = meta["chunks"][i]["file"]
        if seen.get(f, 0) >= 2:
            continue
        seen[f] = seen.get(f, 0) + 1
        out.append({"file": f, "text": meta["chunks"][i]["text"], "score": round(float(s[i]), 3)})
        if len(out) == k:
            break
    return out


# ---------- LLM tool ----------
TOOL = {"type": "function", "function": {
    "name": "search_docs",
    # qwen3:8b ignored a neutral description and answered "I don't have access" (measured); it needs the explicit trigger list
    "description": "Semantic search (Thai/English) over the user's Obsidian vault: long-term memory notes plus their own notes and documents "
                   "(md, txt, PDF, Word). Call it for any question about the user's life or things (people, projects, health, plans, "
                   "passwords, appointments, travel...) when the notes already shown are not enough. Never say you have no access before searching. "
                   "Returns the most relevant passages with file names.",
    "parameters": {"type": "object", "properties": {"query": {"type": "string"}}, "required": ["query"]}}}


def search_docs(query):
    hits = search(query, 3)
    if not hits:
        return "No matching documents found."
    return "\n\n".join(f"[{h['file']}] (score {h['score']})\n{h['text']}" for h in hits)


if __name__ == "__main__":  # selftest: python docsearch.py
    import shutil, tempfile
    c = chunk("ก" * 1000)
    assert len(c) == 4 and all(len(x) <= CHUNK for x in c), [len(x) for x in c]
    assert chunk("hello\nworld") == ["hello\nworld"] and chunk("") == []
    d = tempfile.mkdtemp()
    docs, idx = os.path.join(d, "docs"), os.path.join(d, "docs", ".index")
    os.makedirs(docs)
    with open(os.path.join(d, "old.txt"), "wb") as f:
        f.write("ทดสอบ".encode("cp874"))
    assert read_text(os.path.join(d, "old.txt")) == "ทดสอบ"
    with open(os.path.join(docs, "wifi.txt"), "w", encoding="utf-8") as f:
        f.write("รหัสผ่านไวไฟที่บ้านคือ Somchai2024 เราเตอร์อยู่ห้องนั่งเล่น")
    with open(os.path.join(docs, "plants.md"), "w", encoding="utf-8") as f:
        f.write("# Plants\nWater the monstera once a week; the fiddle leaf fig every 10 days.")
    s1 = build_index(docs, idx)
    assert s1["files"] == 2 and s1["embedded"] == 2, s1
    assert search("รหัสไวไฟบ้าน", 1, idx)[0]["file"] == "wifi.txt"
    assert search("how often to water monstera", 1, idx)[0]["file"] == "plants.md"
    assert build_index(docs, idx)["embedded"] == 0  # nothing changed -> nothing re-embedded
    os.remove(os.path.join(docs, "plants.md"))
    s3 = build_index(docs, idx)
    assert s3["files"] == 1 and s3["chunks"] == 1, s3
    assert build_index(docs, idx, skip=lambda p: p.endswith("wifi.txt"))["files"] == 0 and search("ไวไฟ", 1, idx) == []  # deny drops it
    shutil.rmtree(d)
    print("SELFTEST OK")
