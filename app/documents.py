"""Local documents for the skills that replace Claude Code's docx/pdf/xlsx/pptx ones (written for FRIDAY, nothing copied):
find by file name, read text out (read-only), write NEW files into Documents\\FRIDAY (never overwrite), merge/split PDFs.
Permission checks and confirm() live in friday.py; this module only touches files it is told to."""
import os, re

READ_EXT = (".pdf", ".docx", ".xlsx", ".pptx", ".txt", ".md", ".csv")
READ_MAX, FILE_MAX = 4000, 50 * 2**20  # 4B model, 8k context: one document page-ish at a time
OUT_DIR = os.environ.get("FRIDAY_OUT", os.path.join(os.path.expanduser("~"), "Documents", "FRIDAY"))


def doc_dirs():
    home = os.path.expanduser("~")
    import memory
    return [os.path.join(home, d) for d in ("Documents", "Desktop", "Downloads", os.path.join("OneDrive", "Documents"))] + [memory.VAULT]


def find(query, exts, dirs, limit=8):
    """Newest files with these extensions whose name has every query word; no match = newest ones (with a note)."""
    found = []
    for top in dirs:
        for root, subdirs, files in os.walk(top):
            if root[len(top):].count(os.sep) >= 3:  # ponytail: depth cap instead of an index; deep trees are skipped
                subdirs.clear()
            subdirs[:] = [d for d in subdirs if not d.startswith(".")]
            found += [os.path.join(root, f) for f in files if f.lower().endswith(exts) and not f.startswith("~$")]
    if not found:
        return "No matching files found."
    words = [w for w in re.split(r"\W+", query.lower()) if w]
    hits = [p for p in found if words and all(w in os.path.basename(p).lower() for w in words)]
    (hits or found).sort(key=_mtime, reverse=True)
    note = "" if hits or not words else f"No file name contains '{query}'. Newest files instead:\n"
    return note + "\n".join((hits or found)[:limit])


def _mtime(p):
    try:
        return os.path.getmtime(p)
    except OSError:  # deleted/renamed since the walk (a download finishing), or too long a path
        return 0


_cache = {}  # path -> (mtime, text): reading part 2 of a long PDF must not parse all of it again


def read(path, start=0):
    """Text of a document, READ_MAX chars from `start` (say so when there is more)."""
    path = str(path).strip().strip('"')
    if not path.lower().endswith(READ_EXT):
        return f"Can only read {', '.join(READ_EXT)} files."
    if not os.path.isfile(path):
        return f"No file {path}. Use find_documents first."
    if os.path.getsize(path) > FILE_MAX:
        return "File is larger than 50 MB."
    ext = os.path.splitext(path)[1].lower()
    key = (os.path.realpath(path), _mtime(path))
    if key in _cache:
        text = _cache[key]
    elif ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:  # read_only keeps the file open until close(): Excel could not save it
            text = "\n\n".join(f"## {ws.title}\n" + "\n".join(" | ".join("" if v is None else str(v) for v in row)
                                                              for row in ws.iter_rows(values_only=True) if any(v is not None for v in row))
                               for ws in wb.worksheets)
        finally:
            wb.close()
    elif ext == ".pptx":
        from pptx import Presentation
        text = "\n\n".join(f"## Slide {i}\n" + "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
                           for i, slide in enumerate(Presentation(path).slides, 1))
    else:
        import docsearch
        text = docsearch.read_text(path)
    _cache.clear()
    _cache[key] = text
    start = max(0, int(start or 0))
    part = text[start:start + READ_MAX]
    more = len(text) - start - len(part)
    return part + (f"\n\n[{more} more characters: call read_document again with start={start + len(part)}]" if more > 0 else "")


def _new_path(title, ext):
    os.makedirs(OUT_DIR, exist_ok=True)
    base = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', " ", str(title)).strip(" .")[:80] or "FRIDAY"
    path, n = os.path.join(OUT_DIR, base + ext), 2
    while os.path.exists(path):  # never overwrite: "name (2).docx"
        path, n = os.path.join(OUT_DIR, f"{base} ({n}){ext}"), n + 1
    return path


def _rows(content):
    """Markdown table or CSV lines -> rows (separator lines dropped)."""
    rows = []
    for line in str(content).splitlines():
        line = line.strip()
        if not line or re.fullmatch(r"\|?[\s:|-]+\|?", line):
            continue
        cells = line.strip("|").split("|") if "|" in line else line.split(",")
        rows.append([c.strip() for c in cells])
    return rows


def write(kind, title, content):
    """New file in OUT_DIR. docx/md: '# ' headings, '- ' bullets. xlsx/csv: Markdown table or CSV.
    pptx: each '# ' line starts a slide, lines under it are bullets. Returns the path."""
    kind = str(kind).lower().lstrip(".")
    if kind not in ("docx", "xlsx", "pptx", "md", "csv"):
        raise ValueError("kind must be docx, xlsx, pptx, md or csv")
    path = _new_path(title, "." + kind)
    lines = str(content).splitlines()
    if kind == "md":
        open(path, "w", encoding="utf-8").write(str(content))
    elif kind == "csv":
        import csv
        with open(path, "w", encoding="utf-8-sig", newline="") as f:  # BOM: Excel shows Thai correctly
            csv.writer(f).writerows(_rows(content))
    elif kind == "docx":
        import docx
        d = docx.Document()
        for l in lines:
            m = re.match(r"(#{1,3})\s+(.*)", l)
            if m:
                d.add_heading(m[2], level=len(m[1]))
            elif re.match(r"\s*[-*]\s+", l):
                d.add_paragraph(re.sub(r"^\s*[-*]\s+", "", l), style="List Bullet")
            elif l.strip():
                d.add_paragraph(l.strip())
        d.save(path)
    elif kind == "xlsx":
        import openpyxl
        wb = openpyxl.Workbook()
        for r in _rows(content):
            wb.active.append([float(c) if re.fullmatch(r"-?(0|[1-9]\d{0,14})(\.\d+)?", c) else c for c in r])  # '0812…' stays text
        wb.save(path)
    else:
        from pptx import Presentation
        prs = Presentation()
        slide = None
        for l in lines:
            m = re.match(r"#{1,3}\s+(.*)", l)
            if m or slide is None:
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = m[1] if m else str(title)
                body = slide.placeholders[1].text_frame
                body.text = ""
                if m:
                    continue
            if l.strip():
                p = body.add_paragraph() if body.text else body.paragraphs[0]
                p.text = re.sub(r"^\s*[-*]\s+", "", l).strip()
        prs.save(path)
    return path


def _pages(spec, n):
    out = []
    for part in str(spec).replace(" ", "").split(","):
        a, _, b = part.partition("-")
        if a.isdigit():
            out += range(int(a), (int(b) if b.isdigit() else int(a)) + 1)
    return [p - 1 for p in out if 1 <= p <= n]


def pdf(action, files, pages="", title=""):
    """merge: files (list) -> one new PDF. extract: files[0] + pages '1-3,5' -> new PDF. Returns the path."""
    from pypdf import PdfReader, PdfWriter
    files = [files] if isinstance(files, str) else list(files)
    for f in files:
        if not (str(f).lower().endswith(".pdf") and os.path.isfile(f)):
            raise ValueError(f"not a PDF file: {f}")
    w = PdfWriter()
    if action == "merge":
        for f in files:
            for p in PdfReader(f).pages:
                w.add_page(p)
    elif action == "extract":
        r = PdfReader(files[0])
        keep = _pages(pages, len(r.pages))
        if not keep:
            raise ValueError(f"no valid pages in '{pages}' (file has {len(r.pages)})")
        for i in keep:
            w.add_page(r.pages[i])
    else:
        raise ValueError("action must be merge or extract")
    path = _new_path(title or f"{action} {os.path.splitext(os.path.basename(files[0]))[0]}", ".pdf")
    with open(path, "wb") as f:
        w.write(f)
    return path


if __name__ == "__main__":
    import sys, tempfile
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    OUT_DIR = tempfile.mkdtemp()
    md = "# รายงาน\nบรรทัดแรก\n- ข้อหนึ่ง\n- two"
    tbl = "| ชื่อ | คะแนน | โทร |\n|---|---|---|\n| มิ้นท์ | 9 | 0812345678 |\n| Ann | 7.5 | 1234567890123456 |"
    p_docx, p_docx2 = write("docx", "รายงาน", md), write("docx", "รายงาน", md)
    assert p_docx != p_docx2 and p_docx2.endswith("รายงาน (2).docx")                      # never overwrites
    assert "ข้อหนึ่ง" in read(p_docx) and "รายงาน" in read(p_docx)
    p_x = write("xlsx", "scores", tbl)
    assert "มิ้นท์ | 9 | 0812345678" in read(p_x) and "Ann | 7.5 | 1234567890123456" in read(p_x) and "---" not in read(p_x)
    os.rename(p_x, p_x + ".moved.xlsx")  # workbook closed after reading: the file is not locked
    p_x += ".moved.xlsx"
    p_p = write("pptx", "deck", "# หน้าแรก\n- จุดที่ 1\n- point 2\n# Second\nbody")
    t = read(p_p)
    assert "Slide 1" in t and "จุดที่ 1" in t and "Slide 2" in t and "Second" in t, t
    p_c = write("csv", "t", tbl)
    assert read(p_c).splitlines()[0] == "ชื่อ,คะแนน,โทร"
    assert write("md", "a/b:c*", "x").endswith("a b c.md")                                # no path tricks in titles
    long = write("md", "long", "ก" * 9000)
    assert read(long).endswith("start=4000]") and len(read(long, 8000).split("\n")[0]) == 1000
    assert read(p_docx[:-5] + ".exe").startswith("Can only") and read(os.path.join(OUT_DIR, "no.pdf")).startswith("No file")
    # PDFs: make a 3-page one, extract 2 pages, merge
    from pypdf import PdfWriter, PdfReader
    w = PdfWriter()
    for _ in range(3):
        w.add_blank_page(100, 100)
    src = os.path.join(OUT_DIR, "src.pdf")
    w.write(open(src, "wb"))
    assert len(PdfReader(pdf("extract", [src], "1,3-9")).pages) == 2
    assert len(PdfReader(pdf("merge", [src, src])).pages) == 6
    for bad in (lambda: pdf("extract", [src], "7"), lambda: pdf("merge", [p_docx]), lambda: write("exe", "x", "y")):
        try:
            bad()
            raise AssertionError("should fail")
        except ValueError:
            pass
    r = find("รายงาน", (".docx",), [OUT_DIR]).splitlines()
    assert len(r) == 2 and all(x.endswith(".docx") for x in r), r
    assert find("zzz", READ_EXT, [OUT_DIR]).startswith("No file name")
    print("DOCUMENTS SELFTEST OK")
